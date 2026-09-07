#!/usr/bin/env python3
"""
generate_pptx.py
A helper script that converts the slides_content.md file into a .pptx using python-pptx.
Run this locally after cloning the repository to produce the final PowerPoint file.

Requirements:
  pip install python-pptx markdown

Usage:
  python generate_pptx.py

This script was included because the assistant cannot directly create binary .pptx in the repository via the chat interface. Run locally to generate the final PPTX named 'Buddhist_5A_Analysis.pptx'.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
import textwrap

INPUT = 'slides_content.md'
OUTPUT = 'Buddhist_5A_Analysis.pptx'

prs = Presentation()
prs.slide_height = Inches(11.69)  # A4 height in inches
prs.slide_width = Inches(8.27)    # A4 width in inches

with open(INPUT, 'r', encoding='utf-8') as f:
    content = f.read()

slides = [s.strip() for s in content.split('\n\nSlide: ') if s.strip()]

for s in slides:
    lines = s.split('\n\n', 1)
    title = lines[0].strip()
    body = lines[1].strip() if len(lines) > 1 else ''
    slide_layout = prs.slide_layouts[1]  # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title
    txBox = slide.shapes.placeholders[1].text_frame
    wrapper = textwrap.fill(body, width=90)
    txBox.text = wrapper
    p = txBox.paragraphs[0]
    p.font.size = Pt(12)

prs.save(OUTPUT)
print('Saved:', OUTPUT)
